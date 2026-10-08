#!/usr/bin/env python3
"""Validate FAITHPATH_IMPORT before any writes; materialize approved text verbatim.
Requires requirements-nt-import.txt. Repeat with --check for a read-only integrity check.
"""
import argparse,copy,hashlib,json,re,shutil,sys,unicodedata
from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import load_workbook
HEADERS=['FP_ID','book_key','book_name','chapter','question_order','question','option_0','option_1','option_2','option_3','correct_index','correct_answer','reference','explanation','qa_status','approval','import_ready']
NAME='FaithPath_Kapitelfragen_NT_FINAL_MASTER_1300.xlsx'
BASE='f02106e9a15877ec8acb92f444bf62fcbc51c36e'
def sha(raw):return hashlib.sha256(raw).hexdigest()
def dump(value):return (json.dumps(value,ensure_ascii=False,indent=2)+'\n').encode()
def load(path):return json.loads(path.read_bytes())
def validate_rows(rows,books):
    errors=[];group=defaultdict(list);nt={b['code']:b for b in books[39:]}
    if len(rows)!=1300:errors.append('Datensätze müssen exakt 1300 sein')
    ids=[r.get('FP_ID') for r in rows]
    if len(set(ids))!=1300 or set(ids)!={f'FP-{i:04}' for i in range(261,1561)}:errors.append('FP-ID-Bereich/Einzigartigkeit verletzt')
    for row_no,r in enumerate(rows,2):
        label=f'Excel-Zeile {row_no}, {r.get("FP_ID")}'
        for k in HEADERS:
            if k not in ['chapter','question_order','correct_index'] and (not isinstance(r.get(k),str) or not r[k].strip()):errors.append(f'{label}: {k} fehlt')
        for k in ['chapter','question_order','correct_index']:
            if type(r.get(k)) is not int:errors.append(f'{label}: {k} muss ganzzahlig sein')
        code=r.get('book_key');b=nt.get(code)
        if not b or r.get('book_name')!=b['name'] or type(r.get('chapter')) is not int or not 1<=r['chapter']<=b['chapters']:errors.append(f'{label}: Buch/Kapitel ungültig')
        group[code,r.get('chapter')].append(r)
        options=[r.get(f'option_{i}') for i in range(4)]
        if not all(isinstance(x,str) and x.strip() for x in options) or len(set(options))!=4:errors.append(f'{label}: vier eindeutige Optionen erforderlich')
        c=r.get('correct_index')
        if type(c) is not int or c not in range(4):errors.append(f'{label}: correct_index außerhalb 0–3')
        elif options[c]!=r.get('correct_answer'):errors.append(f'{label}: correct_answer entspricht nicht option_{c}')
        for key,expected in [('qa_status','FREIGEGEBEN'),('approval','FREIGEGEBEN'),('import_ready','JA')]:
            if r.get(key)!=expected:errors.append(f'{label}: {key} muss {expected} sein')
        for key,value in r.items():
            if isinstance(value,str) and ('\ufffd' in value or re.search(r'_x[0-9A-Fa-f]{4}_|&(?:[A-Za-z]+|#\d+|#x[0-9A-Fa-f]+);',value) or any(ord(c)<32 and c not in '\n\t\r' for c in value)):errors.append(f'{label}: Unicode/Excel/Entity-Artefakt in {key}')
    expected={(b['code'],ch) for b in nt.values() for ch in range(1,b['chapters']+1)}
    if set(group)!=expected or len(group)!=260:errors.append('260 vollständige NT-Kapitel erforderlich')
    for key,items in group.items():
        if len(items)!=5 or sorted((r.get('question_order') for r in items),key=str)!=[1,2,3,4,5]:errors.append(f'{key}: genau Frage 1–5 erforderlich')
        questions=[unicodedata.normalize('NFC',r['question']).strip().casefold() for r in items if isinstance(r.get('question'),str)]
        if len(set(questions))!=len(questions):errors.append(f'{key}: doppelte Frage innerhalb Kapitel')
    if errors:raise ValueError('IMPORT STOPP\n'+'\n'.join(errors))
    return group

def parse_sheet(workbook,books):
    # The handoff sheet uses formulas linked to FRAGEN_MASTER. Consume the saved
    # Excel values, never evaluate formulas; absent cached values fail validation.
    w=load_workbook(workbook,read_only=True,data_only=True)
    if 'FAITHPATH_IMPORT' not in w.sheetnames:raise ValueError('IMPORT STOPP: Blatt FAITHPATH_IMPORT fehlt')
    cells=w['FAITHPATH_IMPORT'].iter_rows();header=next(cells)
    if [c.value for c in header]!=HEADERS:raise ValueError('IMPORT STOPP: Headers stimmen nicht mit normalisierter Übergabe überein')
    rows=[]
    for cells in cells:
        if not any(c.value is not None for c in cells):continue
        rows.append(dict(zip(HEADERS,[c.value for c in cells])))
    w.close();validate_rows(rows,books)
    return rows

def validate_references(rows,root,books):
    """Validate all source ranges without editing reference strings."""
    definitions={book['code']:book for book in books}
    chapters={}
    for row in rows:
        reference=row['reference'];prefix=row['book_name']+' '
        if not reference.startswith(prefix):raise ValueError('IMPORT STOPP: Fremdes Buch in Referenz '+row['FP_ID'])
        part=reference[len(prefix):]
        if ',' in part:
            match=re.fullmatch(r'(\d+),(.+)',part)
            if not match:raise ValueError('IMPORT STOPP: Ungültige Referenz '+row['FP_ID'])
            chapter=int(match[1]);part=match[2]
        elif definitions[row['book_key']]['chapters']==1:chapter=1
        else:raise ValueError('IMPORT STOPP: Kapitel fehlt in Referenz '+row['FP_ID'])
        if chapter!=row['chapter']:raise ValueError('IMPORT STOPP: Fremdes Kapitel in Referenz '+row['FP_ID'])
        for segment in re.split(r'[;.]',part):
            match=re.fullmatch(r'(\d+)(?:[–-](\d+))?',segment.strip())
            if not match:raise ValueError('IMPORT STOPP: Ungültiger Versbereich '+row['FP_ID'])
            start=int(match[1]);end=int(match[2] or match[1])
            if start<1 or end<start:raise ValueError('IMPORT STOPP: Ungültiger Versbereich '+row['FP_ID'])
            for translation in ['otb','l1912']:
                key=(translation,row['book_key'])
                if key not in chapters:chapters[key]=load(root/f'data/bibles/{translation}/{row["book_key"]}.json')['chapters']
                present={verse[0] for verse in chapters[key][chapter-1]}
                # Compare bounds before expanding a potentially maliciously huge range.
                if end>max(present) or not all(verse in present for verse in range(start,end+1)):
                    raise ValueError(f'IMPORT STOPP: Vers fehlt {row["FP_ID"]} ({translation}): {reference}')

def import_master(root,workbook,check=False):
    directory=root/'reports/nt-quiz-final';books=load(root/'data/index.json')
    rows=parse_sheet(workbook,books) # The mandated gate is before backup or data writes.
    validate_references(rows,root,books)
    backup=directory/'pre-import-stories.json';original=backup.read_bytes() if backup.exists() else (root/'data/stories.json').read_bytes()
    baseline=load(backup) if backup.exists() else json.loads(original);units=copy.deepcopy(baseline)
    positions={q['faithpath_id']:(si,qi) for si,u in enumerate(units) for qi,q in enumerate(u['questions'])}
    changes=[]
    for r in rows:
        if r['FP_ID'] not in positions:raise ValueError('IMPORT STOPP: Kein bestehender Ort für '+r['FP_ID'])
        si,qi=positions[r['FP_ID']];u=units[si];q=u['questions'][qi]
        if u.get('kind')!='chapter-quiz' or u['book']!=r['book_key'] or u['chapter']+1!=r['chapter'] or qi+1!=r['question_order']:raise ValueError('IMPORT STOPP: ID/Ort passt nicht zu Kapitel '+r['FP_ID'])
        fields={'q':r['question'],'a':[r[f'option_{i}'] for i in range(4)],'c':r['correct_index'],'ref':r['reference'],'x':r['explanation']}
        for key,value in fields.items():
            if q[key]!=value:changes.append({'FP_ID':r['FP_ID'],'locator':f'$[{si}].questions[{qi}].{key}','field':key,'old_value':q[key],'new_value':value})
            q[key]=value
    normalized={'schema_version':1,'sheet':'FAITHPATH_IMPORT','rows':rows}
    result=dump(units);master=dump(normalized)
    status_backup=directory/'pre-import-nt-quiz-status.json'
    old_status=status_backup.read_bytes() if status_backup.exists() else (root/'data/nt-quiz-status.json').read_bytes()
    status=json.loads(old_status)
    status.update(version='FaithPath NT Final Master 1300', method='Imported verbatim from the approved FAITHPATH_IMPORT sheet; zero-based answer indices and existing chapter unit IDs preserved.', note='All 1,300 NT chapter questions are editorially approved: QA_Status and approval FREIGEGEBEN; Import_Bereit JA. Existing 260 curated story questions are unchanged.', source_file=NAME, source_sheet='FAITHPATH_IMPORT', source_sha256=sha(workbook.read_bytes()), qa_status='FREIGEGEBEN', approval='FREIGEGEBEN', import_ready='JA')
    status_bytes=dump(status)
    if backup.exists() and (root/'data/stories.json').read_bytes() not in (original,result):
        raise ValueError('IMPORT STOPP: Bestehende Daten weichen vom gesicherten oder finalen Stand ab; kein Überschreiben fremder Änderungen')
    manifest={'schema_version':1,'source_file':NAME,'source_sheet':'FAITHPATH_IMPORT','source_sha256':sha(workbook.read_bytes()),'base_branch':'develop/v4','base_commit':BASE,'pre_import_sha256':sha(original),'normalized_sha256':sha(master),'imported_stories_sha256':sha(result),'pre_import_status_sha256':sha(old_status),'imported_status_sha256':sha(status_bytes),'questions':1300,'chapters':260,'books':27,'id_start':'FP-0261','id_end':'FP-1560','changed_fields':len(changes),'correct_index_distribution_by_book':{b['code']:dict(sorted(Counter(r['correct_index'] for r in rows if r['book_key']==b['code']).items())) for b in books[39:]}}
    outputs={root/'data/nt-quiz-status.json':status_bytes,directory/'master.json':master,directory/'manifest.json':dump(manifest),directory/'diff.json':dump({'source_sha256':manifest['source_sha256'],'changes':changes}),root/'data/stories.json':result}
    if check:
        for p,raw in outputs.items():
            if not p.exists() or p.read_bytes()!=raw:raise ValueError('IMPORT STOPP: Generiertes Artefakt weicht ab: '+str(p))
    else:
        directory.mkdir(parents=True,exist_ok=True)
        if not backup.exists():backup.write_bytes(original)
        if not status_backup.exists():status_backup.write_bytes(old_status)
        source=directory/NAME
        if workbook.resolve()!=source.resolve():shutil.copyfile(workbook,source)
        for p,raw in outputs.items():
            temp=p.with_name(p.name+'.tmp');temp.write_bytes(raw);temp.replace(p)
    return {'status':'PASS','questions':1300,'chapters':260,'books':27,'changed_fields':len(changes),'mode':'check' if check else 'import'}

def main():
    args=argparse.ArgumentParser();args.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);args.add_argument('--workbook',type=Path);args.add_argument('--check',action='store_true');a=args.parse_args()
    workbook=a.workbook or a.root/'reports/nt-quiz-final'/NAME
    try:print(json.dumps(import_master(a.root,workbook,a.check),ensure_ascii=False))
    except Exception as e:print(str(e),file=sys.stderr);return 1
    return 0
if __name__=='__main__':sys.exit(main())
