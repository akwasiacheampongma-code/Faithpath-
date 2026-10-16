#!/usr/bin/env python3
"""Read-only check of the supplied draft ZIP, Excel workbook and installed baseline."""
import argparse, hashlib, json, zipfile
from pathlib import Path
from openpyxl import load_workbook

def check(root):
    folder = root / 'reports/nachlese-editorial-56'
    source = folder / 'source'
    sha = lambda data: hashlib.sha256(data).hexdigest()
    assert sha((folder/'baseline-mc-questions.json').read_bytes()) == 'f752614dea5cab6c787d3834f3651eb77126cd9c819a998ea4373bc5301d6e37'
    assert sha((source/'Redaktionspaket.zip').read_bytes()) == 'b16804e1dfd7cb8b6e8c48151d72f31b25dae80ca9d969c20b107765457c2eea'
    assert sha((source/'Redaktion_56.xlsx').read_bytes()) == '55408eae434e59d781daad5c80d18b68414cd73293940ff1f8dbe50e3186f681'
    with zipfile.ZipFile(source/'Redaktionspaket.zip') as archive:
        digests = json.loads(archive.read('SHA256.json'))
        for name, digest in digests.items():
            assert Path(name).name == name
            assert sha(archive.read(name)) == digest
            assert (source/name).read_bytes() == archive.read(name)
    proposals = json.loads((source/'56_redaktionsvorschlaege.json').read_text())['records']
    flags = json.loads((source/'mc_optionen_auffaellig.json').read_text())['records']
    baseline = json.loads((folder/'baseline-mc-questions.json').read_text())['questions']
    questions = {q['id']: q for q in baseline}
    assert len(proposals) == len({p['id'] for p in proposals}) == 56
    assert {p['id'] for p in proposals} == {q['id'] for q in baseline if 'a' not in q}
    for p in proposals:
        original = questions[p['id']]
        for a,b in [('story_id','source_story_id'),('reference','ref'),('original_question','q'),('original_answer','answer')]:
            assert p[a] == original[b], (p['id'],a)
        assert len(p['proposed_options']) == 4
        assert p['proposed_correct_answer'] == p['proposed_options'][p['correct_index']]
    assert len(flags) == 227 and len({f['id'] for f in flags}) == 150
    for f in flags:
        q = questions[f['id']]
        assert f['correct'] == q['answer'] and f['suspect_wrong'] in q['a'] and f['suspect_wrong'] != q['answer']
    workbook = load_workbook(source/'Redaktion_56.xlsx',read_only=True,data_only=False,keep_links=False)
    assert workbook.sheetnames == ['56 Entscheidungen','MC Auffaelligkeiten','Übersicht']
    for sheet in workbook:
        assert not any(isinstance(cell,str) and cell.startswith('=') for row in sheet.iter_rows(values_only=True) for cell in row)
    rows = list(workbook['56 Entscheidungen'].iter_rows(values_only=True))
    assert len(rows) == 57
    by_id = {p['id']:p for p in proposals}
    for row in rows[1:]:
        p = by_id[row[0]]
        assert list(row[:11]) == [p['id'],p['story_id'],p['reference'],p['original_question'],p['original_answer'],p['proposed_question'],*p['proposed_options'],p['correct_index']],p['id']
    rows = list(workbook['MC Auffaelligkeiten'].iter_rows(values_only=True))
    assert len(rows) == 228
    for row,flag in zip(rows[1:],flags):
        assert list(row[:5]) == [flag['id'],flag['correct'],flag['suspect_wrong'],flag['similarity'],', '.join(flag['flags'])]
    workbook.close()
    return {'status':'PASS','proposal_rows':56,'heuristic_pairs':227,'marked_questions':150,'excel_json_match':True,'original_fields_match_installed_baseline':True,'input_status':'DRAFT; not a prior human approval','writes':0}

if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check',action='store_true',help='Read-only; this script never writes')
    parser.add_argument('--root',type=Path,default=Path('.'))
    args=parser.parse_args()
    print(json.dumps(check(args.root),ensure_ascii=False))
